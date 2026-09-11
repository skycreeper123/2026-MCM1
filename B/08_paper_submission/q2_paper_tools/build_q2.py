from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile,ZIP_DEFLATED
import hashlib,json
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import doc_helpers as u

HERE=Path(__file__).parent;ROOT=HERE.parents[2]
REF=HERE.parent/'国赛2026全文论文模板.docx';OUT=HERE.parent/'Q2正文.docx'
Q1=HERE.parent/'Q1正文.docx';q1hash=hashlib.sha256(Q1.read_bytes()).hexdigest()
assert hashlib.sha256(REF.read_bytes()).hexdigest()=='f4c3066fe1dfd4b6045f4389d010ac7c0a17933ef15918743cb2e80e34cd8f0a'
reg=json.loads((HERE/'RESULT_REGISTRY.json').read_text(encoding='utf-8'));assert reg['verified']
S=reg['summary'];M=reg['maps'];W=M['wide_fixed'];N=M['narrow_fixed'];I=M['wide_information']
doc=Document(REF);u.doc=doc
sect=deepcopy(doc._element.body.sectPr)
for e in list(doc._element.body):doc._element.body.remove(e)
doc._element.body.append(sect)
para,h,eq,cap,tab=u.para,u.h,u.eq,u.cap,u.tab
def fig(name,title,description):
    p=para('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(HERE/name),width=Cm(16))
    for e in p._p.xpath('.//wp:docPr'):e.set('descr',description)
    cap(title)
def note(s):
    p=para(s,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_before=Pt(4)

h('6 问题二的模型建立与求解',1)
para('针对仅有一次全向干扰源示向观测的情况，本文先构造保证再次收信的第二检测点候选区域，再以第二次观测后定位区域的最坏包围半径为目标选点。通过连续示向响应的区间外包与多起点搜索，兼顾收信可靠性和定位精度，并给出候选区域内定位效果的空间分布。')
h('6.1 首次观测下的源位置外包')
para('记首测点为 S₁，示向度为 θ₁，误差半宽 δ=1°。沿用问题一的前向角域 W(S₁,θ₁,δ)，以 B(c,r) 表示圆心为 c、半径为 r 的闭圆盘。由目标分布范围及最大接收距离，真实源 G 必满足')
eq('G∈F=W(S_{1},θ_{1},δ)∩B(O,1800)∩B(S_{1},1500)',1)
para('正常示向响应还意味着源距大于 5 m。为保持凸性，计算时不剔除这部分近距区域，因而所得集合仍为保守外包。将两个圆分别替换为 K=128 边外切多边形；对圆 B(c,r)，令 uₖ=(cos(2πk/K),sin(2πk/K))，其外包为')
eq(r'Q_{K}(c,r)=⋂_{k=0}^{K−1}\{X:u_{k}·(X−c)≤r\}',2)
para('记两个外切多边形与首测角域的交集为 P⁺，则 F⊆P⁺。外切近似不会漏掉满足题设的真实源；采用问题一的半平面算法得到 P⁺ 的顶点集 V。该物理约束外包用于问题二选点，与问题一仅由示向角域交会得到的定位多边形加以区分。')
h('6.2 保证再次收信的候选区域')
para('全向源的有效接收半径至少为 1000 m。预留 0.1 m 距离裕量，取 ρ=999.9 m，定义固定保证候选区域')
eq(r'C_{safe}=⋂_{v∈V}B(v,ρ)=\{s:max_{v∈V}‖s−v‖_{2}≤ρ\}',3)
para('对任意 G∈P⁺，将 G 表为顶点凸组合，由三角不等式可知，测点到 G 的距离不超过到全部顶点的最大距离。因此式（6-3）保证对全部可能源再次收信。固定候选区是闭圆盘交，具有凸性；其非空条件为')
eq('C_{safe}≠∅ ⇔ R(P^{+})≤ρ,    R(P)=min_{c∈ℝ^{2}}max_{X∈P}‖X−c‖_{2}',4)
para('其中 R(P) 为最小包围圆半径，其圆心提供一个保证收信的可行点。实际选点排除首测点及已测位置。该候选区是保守方案，可超出目标分布圆；处于区外并不意味着无信号。若测点到整个 P⁺ 的最小距离大于 1500 m，则可保守判定为保证无信号，其余位置归为不确定。最小距离计算须考虑多边形内部和边。')

h('6.3 基于连续响应的最坏半径目标',new=True)
h('6.3.1 定位效果指标',3)
para('对候选测点 s，记第二次示向读数为 θ，更新外包为 P(s,θ)=P⁺∩W(s,θ,δ)。根据问题一的覆盖分析，区域直径不能直接代替圆覆盖半径，故采用更新区域的最小包围圆半径作为定位效果指标。记 Θ(s) 为包含所有可能示向读数的角度范围，定义')
eq('H(s)=sup_{θ∈Θ(s)}R(P(s,θ))',5)
para('空交集表示该读数与外包不相容，在取最大值时忽略。H(s) 针对外包模型覆盖所有可能示向响应，具有保守性。实际计算用其上界 U(s) 排名，并可引入移动成本：')
eq(r's^{*}∈arg min_{s∈C_{safe}∖E}J(s),    J(s)=U(s)+λ\frac{‖s−p‖_{2}}{5}',6)
para('其中 E 为已测位置集合，p 为当前位置，默认 p=S₁；速度为 5 m/s，权重 λ 的单位为 m/s。本文主方案取 λ=0，优先改善定位精度；需要兼顾移动时间时可取 λ>0。若收到近距响应，则依据题设距离不超过 5 m 的条件处理，不将其记为几何定位半径等于 0。')
h('6.3.2 示向响应区间外包',3)
para('从 s 指向 V 的方位角构造最短覆盖弧，再向两端各扩展 δ，得到 Θ(s)。若顶点方位不能落入开半圆，或 s 位于顶点附近，则使用完整 360°，从而包含区域内部测点与跨越 0° 的情况。')
para('将读数范围划分为子区间 I=[μ−h,μ+h]。当 δ+h<90° 时，区间内任意角域均被以 μ 为中心、半宽 δ+h 的角域包含，故')
eq('P(s,θ)⊆P^{+}∩W(s,μ,δ+h)=P_{I}(s),    ∀θ∈I',7)
para('由最小包围圆半径对集合包含关系的单调性，各区间外包的半径给出该区间内所有响应半径的上界；当 δ+h≥90° 时，直接采用 P⁺ 作为区间外包。对当前完整划分 𝒯 定义')
eq('U(s)=max_{I∈𝒯}R(P_{I}(s))+η,    L(s)=max_{μ∈𝒜}R(P(s,μ))',8)
para('其中 𝒜 为已计算的区间中点，η=10⁻⁶ m 为数值裕量；空区间不参与最大值。L(s) 是外包模型下的采样下界，不是精确物理源集合上的最坏半径。于是，在几何包含关系成立的前提下，有')
eq('L(s)≤H(s)≤U(s),    Δ(s)=U(s)−L(s)',9)
para('优先二分上界最大的区间，直至上下界间隙达到容差或预算耗尽。仅在两个子区间都计算完成后替换父区间，避免中途停止产生未覆盖的响应区间。该方法具有连续响应的包含关系依据，但采用浮点几何和显式裕量，并非区间算术的形式化认证。')

h('6.4 多起点分级搜索算法',new=True)
para('候选区内的最坏半径场通常呈现多个较优区域。本文采用粗网格、分散起点细化和局部方向搜索，降低单一起点或网格对齐造成的偏差。输入仅为首测位置、示向度及题设参数，真实源位置不参与选点。')
para('首先，在保证区域生成 75 m 粗网格，加入包围圆心、区域中心和首次示向两侧的径向边界特征点。逐点检查收信约束及重复位置，按边界代表性和空间分散性保留至多 64 个粗候选。λ=0 时不为近距离点保留专门名额。')
para('随后，以至多 32 个响应区间粗评全部候选，将上界目标和采样下界目标各前 5 名提升至 128 个区间复评。选择至多 5 个相距至少 75 m 的起点，在其周围 ±75 m 范围内生成 15 m 细网格，每个起点最多保留 49 个候选，再对前列结果按 128、256 个区间分级加密。')
para('最后，从前 2 个结果开展局部方向搜索，依次采用 15、7.5、3.75 m 步长，每级最多移动 2 次，所有试探点均重新检验安全性。对未达精度要求的最终推荐点继续增加区间数，最多至 1024 个，并更新排序。高精度阶段目标间隙为 0.05 m；若达到上限仍未收敛，则同时输出实际间隙。')
para('初始最小包围圆心及其安全邻域提供不依赖源采样的后备测点，解决窄区域内采样失败的问题。预算耗尽时返回已核验的安全结果，并标明超时或未收敛状态。搜索输出为有限候选及有限数值精度下的较优点，不声称连续全局最优；时间预算采用协作式检查，不能硬中断单次几何运算。')
h('6.5 利用首次收信信息扩展候选区域')
para('对于同一全向源，其有效接收半径 r 在两次检测之间保持不变。首次已收到信号意味着 r≥max(1000,‖G−S₁‖₂)，因此可进一步定义信息修正候选区域')
eq(r'C_{1}=⋂_{G∈P^{+}}B(G,max(1000,‖G−S_{1}‖_{2}))',10)
para('固定方案的所有点均属于 C₁。该扩展采用闭集保证，不额外预留固定方案的 0.1 m 裕量。判定 s∈C₁ 等价于验证')
eq('max_{G∈P^{+}}[‖s−G‖_{2}^{2}−max(1000^{2},‖G−S_{1}‖_{2}^{2})]≤0',11)
para('以 S₁ 为圆心的 1000 m 圆内，上式括号内为凸二次函数；圆外为仿射函数，圆周上也可化为仿射表达。因此需检查 P⁺ 顶点、各边与该圆的交点，以及位于 P⁺ 内且背离 s 的圆周极值点。仅检验多边形顶点可能遗漏边或圆弧上的违规点。该有限临界点判据适用于同一接收半径的全向源。')
para('主方案使用固定保证区域，以保留明确距离裕量；C₁ 为扩大可选范围的补充方案。候选区扩大不自动意味着推荐点精度提高，仍须使用同一半径目标进行比较。')

h('6.6 典型候选区域与第二测点结果',new=True)
para('选取两组合成首测观测：大范围案例为 (−1000,0) m、0.5°；边界案例为 (1790,−1400) m、89.3°。固定方案的推荐结果见表 6-1，各指标取自同一轮计算。')
cap('表 6-1 两组典型案例的推荐第二测点',True)
rows=[]
for label,key,unit in [('第二点横坐标','x','m'),('第二点纵坐标','y','m')]:
    rows.append([label+' / '+unit,f'{W["selected"]["position"][key]:.3f}',f'{N["selected"]["position"][key]:.3f}'])
for label,key in [('半径上界 U / m','worst_updated_cover_radius_m'),('上下界间隙 Δ / m','response_bound_gap_m'),('移动距离 / m','movement_distance_m'),('移动时间 / s','movement_time_s')]:
    rows.append([label,f'{W["selected"][key]:.3f}',f'{N["selected"][key]:.3f}'])
tab(['指标','大范围案例','边界案例'],rows,[6.2,4.9,4.9])
note('注：数值保留三位小数用于展示，安全判据使用未舍入坐标；边界点不宜直接按舍入值执行。')
fig('q2_region_paper.png','图 6-1 保证收信区域内的半径上界及数值近优边界','两组候选区域的半径上界热力图，红线为5%数值近优边界，星号为各连通区已评分推荐点。')
note('注：左右色标范围不同；星号为各连通区已评分推荐点，并非全局最优点。')
para('在整个保证区域构建安全三角网格，对 U(s) 作分片线性插值，以采样最小上界的 1.05 倍提取近优轮廓。凸性保证网格三角形位于收信区域内，但不保证三角形内部每一点都满足精度阈值，轮廓、面积与形心均为网格分辨率下的近似。')
para(f'大范围案例的固定候选区面积约 {W["field"]["safe_area_approx_m2"]:,.0f} m²，30 m 网格下近优面积约 {W["field"]["good_area_approx_m2"]:,.0f} m²，分成两个连通区域。45 m 网格所得近优面积约 {M["wide_grid_check"]["field"]["good_area_approx_m2"]:,.0f} m²，说明面积随离散精度变化。采用 C₁ 后候选面积增至约 {I["field"]["safe_area_approx_m2"]:,.0f} m²，但推荐 U 仍约 {I["selected"]["worst_updated_cover_radius_m"]:.3f} m。')

h('6.7 验证结果与适用范围',new=True)
para('采用 5 个种子的 200 个分层随机首测观测及 20 个切向极限案例进行离线合成实验。以原单起点网格为基线，两策略推荐点统一按 256 个响应区间上限和 0.05 m 目标容差复评。')
cap('表 6-2 分层合成实验的配对比较与可靠性检验',True)
tab(['指标','结果'],[
    ['随机组基线与本文半径上界中位数 / m',f'{S["common_old_median_radius_m"]:.3f} / {S["common_new_median_radius_m"]:.3f}'],
    ['随机组配对改善均值与中位数 / m',f'{S["mean_paired_improvement_m"]:.3f} / {S["median_paired_improvement_m"]:.3f}'],
    ['随机组改善 / 持平 / 退步例数','171 / 8 / 21'],
    ['随机组最大单例退步 / m',f'{S["largest_regression_m"]:.3f}'],
    ['220 例失败 / 收信违规 / 真源遗漏 / 上界违规','0 / 0 / 0 / 0'],
    ['基线与本文最大响应上下界间隙 / m',f'{S["old_max_gap_m"]:.3f} / {S["new_max_gap_m"]:.3f}'],
    ['随机组基线与本文计算时间中位数 / s',f'{S["old_median_elapsed_s"]:.3f} / {S["new_median_elapsed_s"]:.3f}']
],[11.3,4.7])
note('注：改善定义为基线 U 减本文 U，按 ±0.05 m 区分改善、持平与退步；时间为原实验记录。')
fig('q2_comparison_paper.png','图 6-2 统一评分精度下的配对改善','200个随机案例的半径上界配对散点图及改善量分布。')
para(f'平均改善的 bootstrap 95% 区间为 [{S["mean_improvement_bootstrap_95ci_m"][0]:.3f},{S["mean_improvement_bootstrap_95ci_m"][1]:.3f}] m，5 个种子均有正的平均改善。结果支持本组实验中的总体改善，但部分案例退步，且计算时间增加。')
para('独立验证检查 49,077 次示向响应，另记 40 次近距响应，未发现上界违规；31 项单元测试全部通过。10 例全域 25 m 网格及局部 5 m 网格复核均通过，最大差距为 0.556 m，验收门槛为 2 m 与参考值 2% 中的较大者。')
para('220 例的最大间隙 0.358 m 满足 0.5 m 总体验收门槛，但未全部达到 0.05 m。圆外包、空间离散和浮点计算仍限制精度，有限网格检验不构成连续全局最优证明。')

doc.core_properties.title='问题二的模型建立与求解'
doc.core_properties.subject='全向干扰源第二检测点选择与保证收信候选区域'
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
temp=HERE/'authored.docx';doc.save(temp)
changed={'word/document.xml','word/_rels/document.xml.rels','docProps/core.xml'}
with ZipFile(REF) as src,ZipFile(temp) as authored,ZipFile(OUT,'w',ZIP_DEFLATED) as dst:
    for n in src.namelist():dst.writestr(n,authored.read(n) if n in changed else src.read(n))
    for n in set(authored.namelist())-set(src.namelist()):dst.writestr(n,authored.read(n))
with ZipFile(REF) as src,ZipFile(OUT) as final:
    preserved=[n for n in src.namelist() if n not in changed]
    assert all(src.read(n)==final.read(n) for n in preserved)
    xml=final.read('word/document.xml').decode('utf-8')
    assert all(t not in xml for t in ['【','待补','but48'])
    assert xml.count('<m:oMath>')==11
assert hashlib.sha256(Q1.read_bytes()).hexdigest()==q1hash
(HERE/'package_qa.json').write_text(json.dumps({'preserved_parts':preserved,'equations':11,'tables':2,'figures':2,'Q1_unchanged_sha256':q1hash,'output_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest()},indent=2),encoding='utf-8')
print(OUT)
