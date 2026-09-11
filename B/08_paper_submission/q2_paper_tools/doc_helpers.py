from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import json, hashlib, re
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc=None
def para(text='',style='Normal'):
    p=doc.add_paragraph(text,style)
    return p
def h(text,level=2,new=False):
    p=para(text,f'Heading {level}')
    p.paragraph_format.page_break_before=new
    return p
def mr(text):
    r=OxmlElement('m:r'); t=OxmlElement('m:t');t.text=text;r.append(t);return r
def wrap(tag,children):
    e=OxmlElement('m:'+tag)
    for x in children:e.append(x)
    return e
def math_parse(s):
    pos=0
    def group():
        nonlocal pos
        if pos<len(s) and s[pos]=='{':pos+=1;return seq('}')
        c=s[pos];pos+=1;return [mr(c)]
    def seq(end=None):
        nonlocal pos
        out=[]
        while pos<len(s):
            c=s[pos]
            if c==end:pos+=1;break
            operator=next((token for token in ['arg min','arg max','safe','sup','max','min','cos','sin'] if s.startswith(token,pos)),None)
            if operator:
                run=mr(operator);pr=OxmlElement('m:rPr');pr.append(OxmlElement('m:nor'));run.insert(0,pr)
                out.append(run);pos+=len(operator);continue
            if s.startswith('\\{',pos) or s.startswith('\\}',pos):
                out.append(mr(s[pos+1]));pos+=2;continue
            if s.startswith('\\frac',pos):
                pos+=5;n=group();d=group(); out.append(wrap('f',[wrap('num',n),wrap('den',d)]));continue
            if c in '_^':
                pos+=1;base=out.pop();g=group()
                if c=='^' and base.tag==qn('m:sSub'):
                    out.append(wrap('sSubSup',[deepcopy(base.find(qn('m:e'))),deepcopy(base.find(qn('m:sub'))),wrap('sup',g)]))
                else:
                    out.append(wrap('sSub' if c=='_' else 'sSup',[wrap('e',[base]),wrap('sub' if c=='_' else 'sup',g)]))
                continue
            if c=='{':pos+=1;out.extend(seq('}'));continue
            out.append(mr(c));pos+=1
        return out
    return seq()
def eq(s,n):
    p=para('','Equation');p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.keep_together=True
    p.add_run('\t');p._p.append(wrap('oMath',math_parse(s)));p.add_run(f'\t（6-{n}）')
def cap(s,before=False):
    p=para(s,'Caption');p.paragraph_format.keep_with_next=before
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
    return p
def tab(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for row in [headers]+rows:
        cells=t.rows[0].cells if row is headers else t.add_row().cells
        for c,w,s in zip(cells,widths,row):c.width=Cm(w);c.text=str(s)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+k),v)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'80' if side in ['top','bottom'] else '100');e.set(qn('w:type'),'dxa');margins.append(e)
    t._tbl.tblPr.append(margins)
    for i,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if i==0:tr.append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                e=OxmlElement('w:shd');e.set(qn('w:fill'),'F2F2F2');c._tc.get_or_add_tcPr().append(e)
            for p in c.paragraphs:
                p.style=doc.styles['Table Header' if i==0 else 'Table Text']
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
    return t

