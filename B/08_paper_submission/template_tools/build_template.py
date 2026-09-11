from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import json

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / 'template_tools' / 'qa'
QA.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
sec.header_distance, sec.footer_distance = Cm(1), Cm(1.25)
sec.different_first_page_header_footer = False

def font(style, east='宋体', size=12, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rp = style.element.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rp.append(rf)
    for k in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        rf.attrib.pop(qn('w:' + k), None)
    rf.set(qn('w:eastAsia'), east)
    rf.set(qn('w:ascii'), 'Times New Roman')
    rf.set(qn('w:hAnsi'), 'Times New Roman')

normal = doc.styles['Normal']
font(normal)
pf = normal.paragraph_format
pf.line_spacing = 1.25
pf.first_line_indent = Pt(24)
pf.space_before = pf.space_after = Pt(0)
pf.widow_control = True
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size, before, after in [('Title',16,0,16),('Heading 1',14,14,7),('Heading 2',12,9,4),('Heading 3',12,6,3)]:
    s = doc.styles[name]; font(s,'黑体',size,True)
    p = s.paragraph_format
    p.first_line_indent = Pt(0); p.line_spacing = 1.15
    p.space_before = Pt(before); p.space_after = Pt(after)
    p.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if name=='Title' else WD_ALIGN_PARAGRAPH.LEFT
    p.page_break_before = False
    borders = s.element.xpath('./w:pPr/w:pBdr')
    for b in borders: b.getparent().remove(b)

for name,size in [('Caption',10.5),('Table Text',10.5),('Table Header',10.5),('Equation',12),('Reference',10.5),('Code',9),('Template Note',10.5),('Unnumbered Heading',14)]:
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = normal
    font(s,'黑体' if name in ('Table Header','Unnumbered Heading') else '宋体',size,name in ('Table Header','Unnumbered Heading'))
    p = s.paragraph_format
    p.first_line_indent=Pt(0); p.line_spacing=1.15
    p.space_before=Pt(0); p.space_after=Pt(4)
    if name in ('Caption','Equation'): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    else: p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if name=='Unnumbered Heading':
        p.space_before=Pt(14);p.space_after=Pt(7);p.keep_with_next=True
    if name=='Reference': p.left_indent=Pt(21);p.first_line_indent=Pt(-21)
    if name=='Code':
        s.font.name='Consolas'
        s.element.rPr.rFonts.set(qn('w:ascii'),'Consolas')
        s.element.rPr.rFonts.set(qn('w:hAnsi'),'Consolas')
        p.line_spacing=1.0;p.space_after=Pt(0)

foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.paragraph_format.first_line_indent=Pt(0)
run=foot.add_run();run.font.size=Pt(10.5)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);field.append(r);foot._p.append(field)
pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'1');sec._sectPr.append(pg)

def p(text, style=None, boldlead=None):
    par=doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        par.add_run(boldlead).bold=True;par.add_run(text[len(boldlead):])
    else: par.add_run(text)
    return par
def h(text,level=1,new=False):
    par=doc.add_paragraph(text,f'Heading {level}')
    if new: par.paragraph_format.page_break_before=True
    return par
def caption(text, before=False):
    par=p(text,'Caption');par.paragraph_format.keep_with_next=before
    par.paragraph_format.space_before=Pt(5)
    return par
def table(headers, rows, widths):
    tab=doc.add_table(rows=1,cols=len(headers))
    tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    for col,w in zip(tab.columns,widths):col.width=Cm(w)
    for c,w,txt in zip(tab.rows[0].cells,widths,headers):c.width=Cm(w);c.text=txt
    for row in rows:
        cells=tab.add_row().cells
        for c,w,txt in zip(cells,widths,row):c.width=Cm(w);c.text=txt
    pr=tab._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    pr.append(borders)
    mar=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        m=OxmlElement('w:'+side);m.set(qn('w:w'),'80' if side in ['top','bottom'] else '100');m.set(qn('w:type'),'dxa');mar.append(m)
    pr.append(mar)
    for i,row in enumerate(tab.rows):
        trpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');c._tc.get_or_add_tcPr().append(shade)
            for par in c.paragraphs:
                par.style=doc.styles['Table Header' if i==0 else 'Table Text']
                par.paragraph_format.space_after=Pt(0)
                par.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or j==0 else WD_ALIGN_PARAGRAPH.LEFT
    p('').paragraph_format.space_after=Pt(0)
    return tab

def equation():
    par=p('','Equation')
    par.paragraph_format.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER)
    par.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
    par.alignment=WD_ALIGN_PARAGRAPH.LEFT
    par.add_run('\t')
    math=OxmlElement('m:oMath')
    mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text='J(x) = f(x) + λg(x)';mr.append(mt);math.append(mr);par._p.append(math)
    par.add_run('\t(1)')

# Page 1: abstract sheet
p('【论文题目】','Title')
par=p('摘  要','Unnumbered Heading');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
par.paragraph_format.space_before=Pt(0)
p('【研究目标】概括题目所研究的对象、主要约束及需要回答的问题。用两至三句话交代建模目标与总体思路，避免大段重复题面或介绍无关背景。')
p('针对问题一，【填写模型名称及关键处理方法】。通过【填写几何推导、计算方法或算法步骤】得到【填写已核验的主要结论及必要数值】，并使用【填写验证方法】检验结论的有效性。',boldlead='针对问题一，')
p('针对问题二，【填写决策变量、优化目标及约束】。建立【填写模型名称】，采用【填写求解方法】获得【填写结果及适用条件】；与【填写有依据的基准方案】比较时，使用一致的评价指标。',boldlead='针对问题二，')
p('针对问题三，【填写策略的主要组成及关键机制】。在【填写真实的实验条件、样本量和测试类型】下，得到【填写已核验结果】；说明方法在效果、耗时或稳健性方面的主要表现。',boldlead='针对问题三，')
p('针对问题四，【填写相对问题三新增的条件和处理方法】。通过【填写检验方法】分析其对【填写评价指标】的影响，给出【填写结论】，并明确方法的适用范围。',boldlead='针对问题四，')
p('【总体结论】用一至两句话概括本文方法的实际价值与主要局限。摘要应使读者能够独立理解问题、方法与结果，不引入正文尚未解释的符号或没有证据支持的性能承诺。')
kw=p('关键词： 【关键词一】；【关键词二】；【关键词三】；【关键词四】')
kw.paragraph_format.first_line_indent=Pt(0);kw.paragraph_format.space_before=Pt(12)
kw.runs[0].bold=True

# Page 2
h('1 问题重述',new=True)
h('1.1 问题背景',2)
p('【填写背景】围绕题目给定的研究对象，介绍与建模直接相关的场景、已知条件及任务目标。保留影响模型建立的信息，避免照抄题面。引用公开资料时在相应位置标注真实文献编号。')
h('1.2 任务分解',2)
p('【填写任务】依次概括四个小问的输入、需要决定的量和输出要求。明确各小问之间共享的信息，以及后续问题相对前述问题新增的约束。')
h('2 问题分析')
p('【填写分析】解释问题的关键难点与可采用的建模路线。围绕信息获取、模型求解与结果验证展开，说明各阶段如何衔接；此处提出方案依据，不提前宣称尚未验证的结论。')
h('3 模型假设与依据')
for i,topic in enumerate(['研究范围与环境条件','测量误差或数据可靠性','算法执行条件与简化处理'],1):
    p(f'假设 {i}：【{topic}】。依据：【说明题面依据或合理性】。影响：【说明适用边界及可能造成的偏差】。')
h('4 符号说明')
caption('表 1 主要符号及其含义',True)
table(['符号','含义','单位或范围'],[
    ['x','【决策变量或状态变量】','【单位或取值域】'],
    ['Ω','【模型中的可行域】','【集合定义】'],
    ['J','【目标函数或评价指标】','【单位】'],
    ['λ','【权重或控制参数】','【无量纲或单位】'],
],[2.1,9.4,4.5])

# Page 3
h('5 问题一的模型建立与求解',new=True)
h('5.1 建模思路与数学表达',2)
p('【填写思路】从已知条件出发定义变量、可行域与目标函数，说明各约束的物理或几何含义。推导过程保留关键环节，并明确每一步成立的前提。下式仅展示可编辑公式与右侧编号样式，应替换为本问实际公式。')
equation()
p('【填写公式解释】逐一说明式（1）中新出现符号的含义、单位和参数来源。若引入某种假设或近似，应解释其合理性及可能影响。不要把示例公式作为已经确定的模型。')
h('5.2 求解方法与结果',2)
p('【填写算法】按照输入、关键步骤、输出和终止条件描述求解过程。说明边界情况及数值误差的处理方式；算法较长时，将完整代码放入附录，正文保留能理解方法的必要步骤。')
p('【填写结果】引用已核验结果，报告必要数值、单位和有效数字。对理论性质给出证明或相应依据，对实验结论交代实验条件，并解释结果如何回答问题一。')
h('6 问题二的模型建立与求解')
h('6.1 决策目标与约束',2)
p('【填写模型】说明问题二与问题一的衔接，定义新增决策变量、优化目标和约束条件。存在多个目标时，交代比较顺序或权重依据；不要将单位不同的指标直接相加。')
h('6.2 算法设计与比较分析',2)
p('【填写求解】描述候选方案的生成、筛选和选择规则，明确无法获得信息或不能满足约束时的处理方案。结合真实结果分析策略选择的理由和适用范围。')
caption('表 2 候选方案的比较结果',True)
table(['方案','主要指标及单位','辅助指标及单位','验证条件'],[
    ['【基准方案】','【待填】','【待填】','【统一条件】'],
    ['【本文方案】','【待填】','【待填】','【统一条件】'],
],[3.2,4.2,4.2,4.4])

# Page 4
h('7 问题三的模型建立与求解',new=True)
h('7.1 总体策略与状态描述',2)
p('【填写策略】定义系统状态、可用观测、允许动作和评价目标，解释各模块之间的数据流与控制逻辑。下图展示流程图的版式；生成论文时，应替换为实际方法图。')
im=Image.new('RGB',(1800,360),'white');draw=ImageDraw.Draw(im)
ff=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',44)
sf=ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',34)
labels=[('信息输入','已知条件与观测'),('模型求解','变量 约束 算法'),('结果检验','复现 比较 误差'),('结论输出','指标与适用范围')]
for i,(a,b) in enumerate(labels):
    x=30+i*450
    draw.rounded_rectangle((x,75,x+380,275),radius=8,outline='black',width=3,fill='#f7f7f7')
    draw.text((x+190,137),a,font=ff,fill='black',anchor='mm')
    draw.text((x+190,209),b,font=sf,fill='black',anchor='mm')
    if i<3:
        draw.line((x+389,175,x+438,175),fill='black',width=3)
        draw.polygon([(x+438,175),(x+426,168),(x+426,182)],fill='black')
im.save(QA/'flow.png')
par=p('');par.paragraph_format.first_line_indent=Pt(0);par.alignment=WD_ALIGN_PARAGRAPH.CENTER
par.paragraph_format.keep_with_next=True
par.add_run().add_picture(str(QA/'flow.png'),width=Cm(16))
caption('图 1 建模与验证流程示意')
h('7.2 关键决策与终止条件',2)
p('【填写机制】说明策略如何选择下一步动作、更新状态、处理失败并判断任务结束。对可能出现的循环、信息缺失或极端条件给出具体处理规则，而不是仅写“进行优化”或“采用智能算法”。')
h('7.3 实验结果与解释',2)
p('【填写实验】清楚区分演练、离线仿真与正式测试，列明实际使用的样本量、参数、版本及指标。用图表呈现关键结果，并说明结论由哪些证据支持。')
h('8 问题四的模型建立与求解')
h('8.1 新增条件与模型调整',2)
p('【填写差异】说明新增条件改变了哪些信息、约束或假设，哪些方法仍可沿用，哪些必须修改。对不同情形分别给出推理，避免直接将问题三的结论推广到全部新场景。')
h('8.2 求解与适用边界',2)
p('【填写结果】描述调整后的算法与验证方式，报告相同口径下的结果。重点分析边界情形、失败案例和方法的局限，说明必要的回退处理及其代价。')

# Page 5
h('9 模型检验与敏感性分析',new=True)
h('9.1 结果复现与有效性检验',2)
p('【填写检验】交代实验条件与复现方法。根据模型特点选择独立计算、解析特例或基准比较，说明检验所支持的结论范围。')
caption('表 3 模型检验与敏感性分析记录',True)
table(['检验项目','设置或变化范围','评价指标','核验结果'],[
    ['基准比较','【统一测试条件】','【指标及单位】','【待填】'],
    ['参数敏感性','【参数及取值范围】','【指标及单位】','【待填】'],
    ['边界情形','【典型极端条件】','【指标及单位】','【待填】'],
],[3,5,4,4])
h('9.2 参数敏感性与稳健性',2)
p('【填写分析】说明关键参数的变化范围及理由，控制其他条件一致，分析结果变化及异常点。没有开展的检验应列为待完成。')
h('10 模型评价与改进')
h('10.1 优点与局限',2)
p('【填写评价】依据正文证据说明模型在解释性、精度或效率方面的优势。同时交代假设限制、适用条件及尚未解决的问题。')
h('10.2 改进方向',2)
p('【填写改进】针对已发现的限制提出可执行的改进方向，区分已经实现的措施与未来工作。不把未开展的实验或设想写成本文成果。')
p('AI工具使用声明','Unnumbered Heading')
p('本参赛队在竞赛过程中使用了AI工具，主要用于【按实际情况填写用途】，详细使用情况见支撑材料。')
p('参考文献','Unnumbered Heading')
p('[1] 【作者】. 【文献题名】[J]. 【期刊名称】, 【年份】, 【卷号】(【期号】): 【页码】.','Reference')
p('[2] 【责任者】. 【网页或报告题名】[EB/OL]. 【发布日期】[【引用日期】]. 【可核验网址】.','Reference')

# Page 6
par=p('附录 A 支撑材料与复现说明','Unnumbered Heading');par.paragraph_format.page_break_before=True
caption('表 A1 支撑材料文件清单',True)
table(['文件或目录','内容与用途','对应位置'],[
    ['【程序文件】','【实现的模型和功能】','【正文小节】'],
    ['【数据文件】','【自主查阅数据及来源】','【正文小节】'],
    ['【结果或日志文件】','【原始输出及复现证据】','【图表编号】'],
    ['AI工具使用详情.pdf','工具、用途、过程及采纳核验情况','AI工具使用声明'],
],[5.1,7.2,3.7])
p('【运行环境】填写语言、依赖版本、必要软件与运行条件。使用相对路径描述输入和输出，不写参赛者身份或本机用户名。')
p('【复现步骤】依次说明准备输入、执行程序、保存结果和核对论文图表的操作。涉及随机算法时记录真实种子；不同小问分别说明入口。')
p('【对应关系】明确每份程序、数据和结果在正文中的用途。支撑包内提供可运行源文件；附录中同时保留全部完整源程序。')
p('附录 B 完整源程序与交互命令','Unnumbered Heading')
p('B.1 【程序文件名】','Heading 2')
p('【此处粘贴完整源程序，使用 Code 样式】','Code')
p('【较长代码允许跨页，保持缩进；不要仅放链接或截图】','Code')
p('【如采用 Excel 或 SPSS 等软件，补齐必要的交互命令】','Code')
p('B.2 【其他程序文件名】','Heading 2')
p('【按支撑材料清单顺序继续收录全部完整源程序】','Code')
p('附录 C 补充推导与结果','Unnumbered Heading')
p('【补充内容】放置有必要但不适合展开在正文中的推导、参数设置或中间结果。正文必须保留主要论证链条，不能把理解模型所必需的内容全部移至附录。')

doc.core_properties.author=''
doc.core_properties.last_modified_by=''
doc.core_properties.title='国赛2026全文论文模板'
doc.core_properties.subject='匿名电子版论文写作模板'
doc.core_properties.comments=''
out=ROOT/'国赛2026全文论文模板.docx'
doc.save(out)

# A plain-text structural companion allows AI to inspect slots without flattening the Word styles.
lines=['# 国赛2026全文生成骨架','', '版式母版：同目录《国赛2026全文论文模板.docx》。本文件只供 AI 读取内容槽位，最终输出应保留 Word 母版样式。', '所有【…】均为待填内容；示例公式和流程图须替换为实际模型。', '']
from docx.text.paragraph import Paragraph
from docx.table import Table
for block in doc.element.body:
    if block.tag == qn('w:p'):
        par=Paragraph(block,doc)
        txt=par.text.strip()
        if block.xpath('.//m:oMath'):
            lines.extend(['【可编辑公式槽位：将示例 J(x) = f(x) + λg(x) 及编号替换为真实公式】',''])
            continue
        if block.xpath('.//w:drawing'):
            lines.extend(['【图形槽位：插入实际方法流程图，沿用母版的居中与图题样式】',''])
            continue
        if not txt: continue
        style=par.style.name
        prefix = '# ' if style=='Title' else ('## ' if style in ('Heading 1','Unnumbered Heading') else ('### ' if style=='Heading 2' else ''))
        lines.extend([prefix+txt,''])
    elif block.tag == qn('w:tbl'):
        tab=Table(block,doc)
        for i,row in enumerate(tab.rows):
            lines.append('| '+' | '.join(c.text for c in row.cells)+' |')
            if i==0: lines.append('| '+' | '.join('---' for c in row.cells)+' |')
        lines.append('')
(ROOT/'国赛2026全文生成骨架.md').write_text('\n'.join(lines),encoding='utf-8')
print(str(out))
print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables))
