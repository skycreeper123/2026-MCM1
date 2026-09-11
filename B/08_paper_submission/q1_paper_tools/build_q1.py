from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import json, hashlib, re
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).parent
ROOT=HERE.parents[2]
REF=HERE.parent/'国赛2026全文论文模板.docx'
OUT=HERE.parent/'Q1正文.docx'
assert hashlib.sha256(REF.read_bytes()).hexdigest()=='f4c3066fe1dfd4b6045f4389d010ac7c0a17933ef15918743cb2e80e34cd8f0a'
reg=json.loads((HERE/'RESULT_REGISTRY.json').read_text(encoding='utf-8'))
assert all(r['verified'] for r in reg['records'])
A,B,R,T=reg['records']; A=A['result']; B=B['result']
doc=Document(REF)
sect=deepcopy(doc._element.body.sectPr)
for e in list(doc._element.body):doc._element.body.remove(e)
doc._element.body.append(sect)

def para(text='',style='Normal'):
    p=doc.add_paragraph(text,style)
    return p
def h(text,level=2,new=False):
    p=para(text,f'Heading {level}')
    p.paragraph_format.page_break_before=new
    return p
def mr(text):
    r=OxmlElement('m:r'); t=OxmlElement('m:t');t.text=text;r.append(t);return r
def wrap(tag,children):
    e=OxmlElement('m:'+tag)
    for x in children:e.append(x)
    return e
def math_parse(s):
    pos=0
    def group():
        nonlocal pos
        if pos<len(s) and s[pos]=='{':pos+=1;return seq('}')
        c=s[pos];pos+=1;return [mr(c)]
    def seq(end=None):
        nonlocal pos
        out=[]
        while pos<len(s):
            c=s[pos]
            if c==end:pos+=1;break
            if s.startswith('\\{',pos) or s.startswith('\\}',pos):
                out.append(mr(s[pos+1]));pos+=2;continue
            if s.startswith('\\frac',pos):
                pos+=5;n=group();d=group(); out.append(wrap('f',[wrap('num',n),wrap('den',d)]));continue
            if c in '_^':
                pos+=1;base=out.pop();g=group()
                if c=='^' and base.tag==qn('m:sSub'):
                    out.append(wrap('sSubSup',[deepcopy(base.find(qn('m:e'))),deepcopy(base.find(qn('m:sub'))),wrap('sup',g)]))
                else:
                    out.append(wrap('sSub' if c=='_' else 'sSup',[wrap('e',[base]),wrap('sub' if c=='_' else 'sup',g)]))
                continue
            if c=='{':pos+=1;out.extend(seq('}'));continue
            out.append(mr(c));pos+=1
        return out
    return seq()
def eq(s,n):
    p=para('','Equation');p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.keep_together=True
    p.add_run('\t');p._p.append(wrap('oMath',math_parse(s)));p.add_run(f'\t（5-{n}）')
def cap(s,before=False):
    p=para(s,'Caption');p.paragraph_format.keep_with_next=before
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
    return p
def tab(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for row in [headers]+rows:
        cells=t.rows[0].cells if row is headers else t.add_row().cells
        for c,w,s in zip(cells,widths,row):c.width=Cm(w);c.text=str(s)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+k),v)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'80' if side in ['top','bottom'] else '100');e.set(qn('w:type'),'dxa');margins.append(e)
    t._tbl.tblPr.append(margins)
    for i,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if i==0:tr.append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                e=OxmlElement('w:shd');e.set(qn('w:fill'),'F2F2F2');c._tc.get_or_add_tcPr().append(e)
            for p in c.paragraphs:
                p.style=doc.styles['Table Header' if i==0 else 'Table Text']
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
    return t

h('5 问题一的模型建立与求解',1)
para('针对示向误差引起的区域定位问题，本文将各检测点的前向角域转化为线性半平面，以半平面交集刻画定位区域，并利用凸性将区域直径和圆盘覆盖检验化为有限顶点计算。分析表明，以最远点对为直径端点的圆不一定覆盖定位区域，必须另行检验各顶点是否位于圆内或圆上。')
h('5.1 交会定位区域的半平面模型')
para('按照题设的平面坐标系，以正东方向为 x 轴正向、正北方向为 y 轴正向，坐标单位为米。设同一干扰源的 n 次有效观测对应检测点 Sᵢ=(xᵢ,yᵢ)，示向度为 θᵢ，从 x 轴正向逆时针计量；误差半宽取 δ=1°。仅使用题设给出的误差界，不对误差分布作额外假设，也不通过同地点重复测量缩小误差界。')
para('示向度确定的是具有前向性的角域。记检测点 i 的下、上边界单位方向向量分别为 lᵢ、rᵢ，则')
eq('l_{i}=(cos(θ_{i}−δ),sin(θ_{i}−δ)),   r_{i}=(cos(θ_{i}+δ),sin(θ_{i}+δ))',1)
para('令待定位点为 X=(x,y)，二维叉积定义为 u×v=uₓvᵧ−uᵧvₓ。由于角域宽度为 2δ<180°，X 位于第 i 个闭角域内的充要条件为')
eq('l_{i}×(X−S_{i})≥0,    r_{i}×(X−S_{i})≤0',2)
para('式（5-2）同时限制边界和方向，可排除沿示向反方向延伸得到的伪交会区域。将叉积展开，可得到两个线性不等式：')
eq('(l_{iy},−l_{ix})·X≤(l_{iy},−l_{ix})·S_{i}',3)
eq('(−r_{iy},r_{ix})·X≤(−r_{iy},r_{ix})·S_{i}',4)
para('汇总 n 次观测形成 2n 个半平面，得到交会定位区域')
eq(r'P=⋂_{i=1}^{n}W_{i}=\{X∈ℝ^{2}:AX≤b\}',5)
para('其中，Wᵢ 为式（5-2）定义的闭角域。半平面均为凸集，因此 P 为闭凸集；当 P 非空且有界时，它是凸多边形，也可能退化为线段或单点。三角函数天然处理跨越 0° 的读数，计算时将角度统一换算为弧度。')
para('本问按附录 2 的交会定位定义求取角域交集，不额外裁入目标分布圆、接收圆或人为方框。由此保留题目要求的多边形定位区域，并避免人为边界改变区域直径。若观测矛盾导致交集为空，或几何布局无法约束有限区域，则分别返回空集、无界状态，而不强行给出有限直径。')

h('5.2 定位区域直径的有限化计算',new=True)
para('以下讨论 P 非空且有界的情形。设其顶点集为 V={v₁,…,vₘ}，区域直径 D 为 P 内任意两点的最大欧氏距离。由于 P 是有限顶点的凸包，任意 X、Y∈P 均可表示为顶点的凸组合，分别记系数为 αᵢ、βⱼ，系数非负且各自之和为 1。由三角不等式，')
eq('‖X−Y‖_{2}=‖∑_{i,j}α_{i}β_{j}(v_{i}−v_{j})‖_{2}≤∑_{i,j}α_{i}β_{j}‖v_{i}−v_{j}‖_{2}',6)
para('右端不超过最大顶点间距；反过来，所有顶点均属于 P，因此该最大间距可以在 P 内取得。故区域直径精确等于')
eq('D=max_{1≤i,j≤m}‖v_{i}−v_{j}‖_{2},    (a,b)∈arg max_{v_{i},v_{j}∈V}‖v_{i}−v_{j}‖_{2}',7)
para('据此枚举全部顶点对，保存最大距离 D 及对应最远点对 a、b，即可完成直径计算。该结论适用于退化区域：线段的直径为两端点距离，单点的直径为 0。')
h('5.3 直径圆的覆盖判据与反例')
para('以最远点对 a、b 为直径端点，构造闭圆盘 C，其圆心和半径分别为 o=(a+b)/2、r=D/2。由于圆盘为凸集，包含全部顶点等价于包含其凸包，因此覆盖判据为')
eq(r'P⊆C ⇔ max_{v∈V}‖v−o‖_{2}≤\frac{D}{2}',8)
para('由平方展开可得等价的点积判据：')
eq(r'(v−a)·(v−b)=‖v−o‖_{2}^{2}−\frac{D^{2}}{4}≤0,    ∀v∈V',9)
para('式（5-9）成立时，所有顶点均位于圆内或圆上；只要有一个顶点对应点积为正，直径圆便不能覆盖定位区域。例如，边长为 2 m 的等边三角形，其区域直径为 2 m，而任一边中点到第三个顶点的距离为 √3 m，大于直径圆半径 1 m。因此，“直径圆总能覆盖定位区域”的命题不成立。该解析反例用于说明一般凸多边形的几何性质；第 5.5.2 节进一步给出满足本题误差约束的测向反例。')
para('更换同半径圆的圆心也不能规避此判据。若某圆心 c、半径 D/2 的圆盘包含 P，则它必须包含 a、b，由三角不等式有')
eq('D=‖a−b‖_{2}≤‖a−c‖_{2}+‖c−b‖_{2}≤D',10)
para('两侧相等迫使 c 位于线段 ab 上且到两端的距离均为 D/2，故 c=o。这说明任取一组最远点对即可判定：若其直径圆不覆盖 P，则不存在其他半径为 D/2 的圆盘覆盖 P。')
para('对于后续半径为 20 m 的光学精确定位，D≤40 m 仅是存在区域覆盖点的必要条件。仍须检验拟选点到全部顶点的最大距离，或进一步求最小包围圆，不能单凭直径阈值保证完成清除。')

h('5.4 求解流程与数值处理',new=True)
para('算法输入为检测点坐标、示向度及误差半宽，输出区域状态、顶点、直径、最远点对和覆盖结论，求解步骤如下。')
para('步骤 1  按式（5-3）和式（5-4）构建 A、b，将法向量归一化，使约束残差具有米制距离意义。用线性规划检验 AX≤b 的可行性，无可行解时返回空集。')
para('步骤 2  分别求解 x、−x、y、−y 的最小化问题，两个坐标均允许取任意实数。任一方向无界则 P 无界，其数学直径为无穷大；仅在四个方向均有界时继续计算。')
para('步骤 3  枚举不平行的边界直线对，求交点并检验全部半平面约束，保留可行交点。按米制容差去重，再取凸包，得到逆时针排列的顶点集 V。以线性规划给出的坐标极值复核顶点是否存在明显遗漏。')
para('步骤 4  根据式（5-7）枚举顶点对，得到 D、a、b，再求圆心 o 和半径 D/2。定义覆盖差值')
eq(r'g=max_{v∈V}‖v−o‖_{2}−\frac{D}{2}',11)
para('计算中取几何容差 ε=10⁻⁷ m，当 g≤ε 时判定覆盖；否则返回不覆盖，并记录圆外顶点。采用距离形式的残差，避免将长度容差直接与式（5-9）中具有面积量纲的点积比较。')
para('交点计算先平移到可行点附近，以减小大坐标引起的数值损失。若线性规划失败、交点异常或极值复核不一致，返回数值异常。直径枚举为 O(m²)，覆盖检验为 O(m)，不含半平面求交等计算，适合本问的小规模区域。')
h('5.5 结果验证与分析')
h('5.5.1 解析检验与随机交叉验证',3)
para('对顶点为 (0,0)、(4,0)、(4,3)、(0,3) 的矩形，计算直径为 5 m，直径圆覆盖；边长为 2 m 的等边三角形直径为 2 m，判定不覆盖，均与解析结论一致。')
para('以固定种子 20260910 生成 40 组三测点观测：真实源在 [−500,500]² m 内均匀采样，测点沿相对真实源的 0°、120°、240° 方向布设，距离在 [100,900] m 内、示向误差在 [−1°,1°] 内均匀采样。该分布仅用于验证。')
para(f'以独立逐边多边形裁剪算法作为求交参照，40 组均得到有界区域并保留真实源。最大直径绝对差为 {R["diameter_error_m"]:.4e} m，最大双向顶点匹配误差为 {R["vertex_error_m"]:.4e} m，均小于 10⁻⁶ m。参照算法使用的初始方框逐例确认未截断最终区域；两实现共用角域转换，因此该对照主要验证求交和直径计算。')
para('18 项单元测试全部通过，覆盖空集、无界、退化、跨零角度及近平行等情形，包含 200 个随机三角形覆盖判据对照和 10 组最小包围圆独立优化对照。上述差异仅反映数值实现的一致性，不是真实源定位误差或任意输入的误差上界。')

h('5.5.2 六测点合成案例与覆盖结论',3,new=True)
para('题面未给出问题一的固定数值观测数据，故选取两组各含六个检测点的合成观测展示算法输出。两组真实源均设为 (250,−150) m，仅用于生成和核验观测，不作为定位算法输入；根据保存后的坐标及示向度重新核算，全部实际角误差均满足 ±1° 约束。两组定位区域均含五个顶点，结果见表 5-1 和图 5-1。')
cap('表 5-1 六测点合成案例的直径与覆盖结果',True)
rows=[]
for name,x in [('A',A),('B',B)]:
    c=x['diameter_circle']
    rows.append([name,f'{x["diameter_m"]:.4f}',f'{c["radius_m"]:.4f}',f'{c["max_vertex_distance_m"]:.4f}',f'{c["coverage_gap_m"]:.4f}','是' if x['diameter_circle_covers'] else '否'])
tab(['案例','直径 D/m','半径 D/2/m','最大顶点距离/m','差值 g/m','覆盖'],rows,[1.1,2.7,2.9,3.5,3.0,1.8])
p=para('注：长度保留四位小数；案例 A 的覆盖差值在数值容差内为零。','Caption');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_before=Pt(4)
p=para('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
p.add_run().add_picture(str(ROOT/'B/q1/validation_results/q1_diameter_coverage.png'),width=Cm(16))
for d in p._p.xpath('.//wp:docPr'):d.set('descr','六测点合成案例A的直径圆覆盖定位多边形，案例B的一个顶点在直径圆外。')
cap('图 5-1 六测点定位区域及其直径圆')
p=para('注：坐标以合成真实源为原点平移，单位为米；两图均为等比例坐标轴，但显示范围不同。浅蓝色为定位区域，红色实线为最远点对，红色虚线为直径圆，橙色为圆外顶点。','Caption');p.alignment=WD_ALIGN_PARAGRAPH.LEFT
para(f'案例 A 的直径为 {A["diameter_m"]:.4f} m，全部顶点均在直径圆内或圆上，因而该圆覆盖整个定位区域。案例 B 的直径为 {B["diameter_m"]:.4f} m，其直径圆半径为 {B["diameter_circle"]["radius_m"]:.4f} m，而最远顶点到圆心的距离为 {B["diameter_circle"]["max_vertex_distance_m"]:.4f} m，超出圆周 {B["diameter_circle"]["coverage_gap_m"]:.4f} m，故不覆盖。该反例由合法误差范围内的测向观测产生，直接回答了本问的覆盖问题。')

doc.core_properties.title='问题一的模型建立与求解'
doc.core_properties.subject='无线电干扰源交会定位区域直径及圆盘覆盖判定'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.comments=''
temp=HERE/'authored.docx';doc.save(temp)
changed={'word/document.xml','word/_rels/document.xml.rels','docProps/core.xml'}
with ZipFile(REF) as src, ZipFile(temp) as authored, ZipFile(OUT,'w',ZIP_DEFLATED) as dst:
    for name in src.namelist():dst.writestr(name,authored.read(name) if name in changed else src.read(name))
    for name in set(authored.namelist())-set(src.namelist()):dst.writestr(name,authored.read(name))
with ZipFile(REF) as src, ZipFile(OUT) as final:
    unchanged={n:hashlib.sha256(src.read(n)).hexdigest() for n in src.namelist() if n not in changed}
    assert all(src.read(n)==final.read(n) for n in unchanged)
    xml=final.read('word/document.xml').decode('utf-8')
    assert '【' not in xml and '[待补' not in xml
    assert xml.count('<m:oMath>')==11
    (HERE/'package_qa.json').write_text(json.dumps({'preserved_parts':unchanged,'editable_parts':sorted(changed),'output_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'equations':11},indent=2),encoding='utf-8')
print(OUT)
